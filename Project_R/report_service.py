"""Generate weekly and monthly XLSX attendance reports."""

import calendar
import re
from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from attendance import get_connection, get_employee, get_employee_attendance_history


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ATTENDANCE_COLUMNS = [
    "Date",
    "Employee ID",
    "Employee Name",
    "Shift Start",
    "Shift End",
    "First In",
    "Last Out",
    "Worked Hours",
    "Worked Minutes",
    "Status",
    "Leave Type",
    "Leave Reason",
]


def _parse_report_date(value):
    if not value:
        return date.today()
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except (TypeError, ValueError) as error:
        raise ValueError("date must use YYYY-MM-DD format") from error


def _period_dates(period, selected_date):
    selected = _parse_report_date(selected_date)
    if period == "weekly":
        start = selected - timedelta(days=selected.weekday())
        return start, start + timedelta(days=6)
    if period == "monthly":
        start = selected.replace(day=1)
        end = selected.replace(day=calendar.monthrange(selected.year, selected.month)[1])
        return start, end
    raise ValueError("period must be weekly or monthly")


def _history_for_range(employee_id, start_date, end_date):
    rows = {}
    cursor = start_date.replace(day=1)
    while cursor <= end_date:
        history = get_employee_attendance_history(employee_id, cursor.year, cursor.month)
        if history is None:
            raise LookupError(f"Employee '{employee_id}' not found")
        rows.update({item["date"]: item for item in history["attendance"]})
        if cursor.month == 12:
            cursor = cursor.replace(year=cursor.year + 1, month=1)
        else:
            cursor = cursor.replace(month=cursor.month + 1)
    return [rows.get((start_date + timedelta(days=offset)).isoformat()) for offset in range((end_date - start_date).days + 1)]


def _employee_rows():
    connection = get_connection()
    try:
        return [dict(row) for row in connection.execute(
            """SELECT employee_id, employee_name, shift_start, shift_end, email
               FROM employees ORDER BY employee_name COLLATE NOCASE"""
        ).fetchall()]
    finally:
        connection.close()


def _safe_filename(value):
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", str(value)).strip("_.") or "employee"


def _style_sheet(sheet, widths):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _add_attendance_rows(workbook, employees, start_date, end_date):
    summary = workbook.active
    summary.title = "Attendance Summary"
    summary.append(ATTENDANCE_COLUMNS)
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    event_sheet = workbook.create_sheet("Attendance Events")
    event_sheet.append(["Date", "Employee ID", "Employee Name", "Event Time", "Event Type"])
    for cell in event_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")

    for employee in employees:
        days = _history_for_range(employee["employee_id"], start_date, end_date)
        for day in days:
            if day is None:
                continue
            summary.append([
                datetime.strptime(day["date"], "%Y-%m-%d").date(),
                employee["employee_id"],
                employee["employee_name"],
                employee.get("shift_start") or "",
                employee.get("shift_end") or "",
                day.get("first_in") or "",
                day.get("last_out") or "",
                day.get("worked_hours") or "0h 0m",
                int(day.get("worked_minutes") or 0),
                day.get("status") or "ABSENT",
                day.get("leave_type") or "",
                day.get("reason") or "",
            ])
            for event in day.get("events", []):
                event_sheet.append([
                    datetime.strptime(day["date"], "%Y-%m-%d").date(),
                    employee["employee_id"],
                    employee["employee_name"],
                    event.get("event_time") or "",
                    event.get("event_type") or "",
                ])

    _style_sheet(summary, [14, 16, 28, 14, 14, 14, 14, 16, 16, 16, 18, 36])
    _style_sheet(event_sheet, [14, 16, 28, 16, 16])


def _create_report(employees, period, selected_date, report_name):
    start_date, end_date = _period_dates(period, selected_date)
    workbook = Workbook()
    _add_attendance_rows(workbook, employees, start_date, end_date)
    output = BytesIO()
    workbook.save(output)
    filename = (
        f"{_safe_filename(report_name)}_{period}_attendance_"
        f"{start_date:%Y%m%d}_to_{end_date:%Y%m%d}.xlsx"
    )
    return output.getvalue(), filename


def create_employee_report(employee_id, period="monthly", selected_date=None):
    employee = get_employee(employee_id)
    if employee is None:
        raise LookupError(f"Employee '{employee_id}' not found")
    return _create_report(
        [dict(employee)], period, selected_date, employee_id
    )


def create_all_employees_report(period="monthly", selected_date=None):
    return _create_report(
        _employee_rows(), period, selected_date, "all_employees"
    )
